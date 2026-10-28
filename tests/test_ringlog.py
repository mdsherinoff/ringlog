import csv
import json
import threading
import urllib.request
from http.server import ThreadingHTTPServer

import pytest
from openpyxl import load_workbook

from fakebackup import make_backup
from ringlog.backup import Backup, PasswordRequired, WrongPassword, find_backups
from ringlog.calls import MISSED, OUTGOING, filter_calls, read_calls
from ringlog.cli import main
from ringlog.export import to_csv, to_html, to_xlsx


@pytest.fixture(params=[None, "s3cret"], ids=["plain", "encrypted"])
def backup_dir(tmp_path, request):
    return make_backup(tmp_path / "ABC123", request.param), request.param


def test_reads_calls_with_contact_names(backup_dir):
    path, pw = backup_dir
    with Backup(path, pw) as b:
        calls = read_calls(b)
    assert len(calls) == 5
    newest = calls[0]
    assert newest.name == "Corner Shop" and newest.service == "WhatsApp"
    # +44 7700 900123 and 07700900123 both match the contact saved as +44 (7700) 900-123
    alex = [c for c in calls if c.name == "Alex Morgan"]
    assert len(alex) == 2
    first = calls[-1]
    assert first.direction == OUTGOING and first.duration == 125 and first.duration_text == "2:05"
    assert [c.direction for c in calls].count(MISSED) == 1
    facetime = next(c for c in calls if c.number == "someone@example.com")
    assert facetime.service == "FaceTime Audio"


def test_encrypted_needs_right_password(tmp_path):
    path = make_backup(tmp_path / "enc", "right")
    with pytest.raises(PasswordRequired):
        Backup(path)
    with pytest.raises(WrongPassword):
        Backup(path, "wrong")


def test_find_backups(tmp_path):
    make_backup(tmp_path / "one")
    make_backup(tmp_path / "two", "pw")
    (tmp_path / "junk").mkdir()
    found = find_backups([tmp_path])
    assert {b.path.name for b in found} == {"one", "two"}
    assert {b.encrypted for b in found} == {True, False}
    assert "Test iPhone" in found[0].label


def test_filter(backup_dir):
    path, pw = backup_dir
    with Backup(path, pw) as b:
        calls = read_calls(b)
    assert len(filter_calls(calls, directions={MISSED})) == 1
    assert len(filter_calls(calls, search="alex")) == 2
    assert len(filter_calls(calls, search="0100 222")) == 1


def test_exports(tmp_path):
    path = make_backup(tmp_path / "b")
    with Backup(path) as b:
        calls = read_calls(b)
    to_csv(calls, tmp_path / "c.csv")
    rows = list(csv.reader(open(tmp_path / "c.csv", encoding="utf-8-sig")))
    assert rows[0][0] == "Date" and len(rows) == 6

    to_xlsx(calls, tmp_path / "c.xlsx")
    wb = load_workbook(tmp_path / "c.xlsx")
    assert wb["Calls"].max_row == 6
    assert wb["Summary"]["B2"].value == 5

    to_html(calls, tmp_path / "r.html")
    page = (tmp_path / "r.html").read_text()
    assert "Most frequent contacts" in page and "Alex Morgan" in page


def test_cli_export(tmp_path, capsys):
    path = make_backup(tmp_path / "b", "pw")
    out = tmp_path / "out.xlsx"
    assert main(["export", str(out), "--backup", str(path), "--password", "pw", "--only", "missed"]) == 0
    assert load_workbook(out)["Calls"].max_row == 2
    assert main(["export", str(tmp_path / "x.txt"), "--backup", str(path), "--password", "pw"]) == 1


def test_web_api(tmp_path):
    from ringlog.web import _make_handler, _State

    path = make_backup(tmp_path / "b", "pw")
    state, ref = _State(), {}
    server = ThreadingHTTPServer(("127.0.0.1", 0), _make_handler(state, ref))
    ref["server"] = server
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_address[1]}"

    def post(route, body, token=state.token):
        req = urllib.request.Request(base + route, json.dumps(body).encode(), {"X-Ringlog-Token": token}, method="POST")
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, r.read()
        except urllib.error.HTTPError as e:
            return e.code, e.read()

    try:
        assert state.token in urllib.request.urlopen(base + "/").read().decode()
        assert post("/api/open", {"path": str(path)}, token="nope")[0] == 403
        status, body = post("/api/open", {"path": str(path)})
        assert status == 401 and json.loads(body)["needPassword"]
        status, body = post("/api/open", {"path": str(path), "password": "pw"})
        assert status == 200 and len(json.loads(body)["calls"]) == 5
        status, body = post("/api/export", {"format": "csv", "directions": ["Missed"]})
        assert status == 200 and body.decode("utf-8-sig").count("\n") == 2
    finally:
        server.shutdown()
